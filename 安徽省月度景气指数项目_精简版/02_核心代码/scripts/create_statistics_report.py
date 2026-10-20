from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data" / "raw"
OUTPUT = ROOT / "data" / "processed" / "安徽省月度景气指数统计报表.xlsx"

CORE_COLUMNS = [
    "工业增加值同比增长 (%)",
    "固定资产投资累计增长 (%)",
    "利润总额累计增长 (%)",
    "发电量同比增长 (%)",
    "PPI同比偏离值",
]

CORE_VARIABLE_LABELS = {
    "月份": "时间",
    "工业增加值同比增长 (%)": "变量1：工业增加值同比（%）",
    "固定资产投资累计增长 (%)": "变量2：固定资产投资累计同比（%）",
    "利润总额累计增长 (%)": "变量3：利润总额累计同比（%）",
    "发电量同比增长 (%)": "变量4：发电量同比（%）",
    "PPI同比偏离值": "变量5：PPI同比偏离值（百分点）",
    "完整核心指标数": "完整核心指标数",
    "质量标记": "质量标记",
}


def build_core_wide_table(wide: pd.DataFrame) -> pd.DataFrame:
    """Return one row per month with the five numbered modelling variables."""
    columns = ["月份", *CORE_COLUMNS, "完整核心指标数", "质量标记"]
    return wide.loc[:, columns].rename(columns=CORE_VARIABLE_LABELS).copy()


def _format_workbook(path: Path) -> None:
    workbook = load_workbook(path)
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, (date, datetime)):
                    cell.number_format = "yyyy-mm-dd"
        sheet.column_dimensions["A"].width = max(
            sheet.column_dimensions["A"].width or 0, 12
        )
    workbook.save(path)


def read_export(path: Path) -> pd.DataFrame:
    sheet = load_workbook(path, read_only=True, data_only=True).active
    rows = list(sheet.iter_rows(values_only=True))
    headers = [str(value).strip() if value is not None else None for value in rows[3]]
    records: list[dict[str, object]] = []
    for row in rows[4:]:
        indicator = str(row[0]).strip() if row and row[0] is not None else ""
        if not indicator or indicator.startswith("注"):
            continue
        for header, value in zip(headers[1:], row[1:]):
            if not header or value in (None, ""):
                continue
            records.append({"月份": pd.to_datetime(header.replace("年", "-").replace("月", "-01")), "指标": indicator, "数值": float(value)})
    return pd.DataFrame(records)


def main() -> None:
    files = sorted(RAW.glob("provincial_monthly_*.xlsx"))
    if not files:
        raise FileNotFoundError("data/raw 中需要完整的安徽省分省月度数据报表")

    long = pd.concat([read_export(path) for path in files], ignore_index=True)
    long = long.drop_duplicates(["月份", "指标"], keep="last").sort_values(["月份", "指标"])
    wide = long.pivot(index="月份", columns="指标", values="数值").sort_index()
    rename = {
        "工业增加值同比增长 (%)": "工业增加值同比增长 (%)",
        "固定资产投资完成额累计增长 (%)": "固定资产投资累计增长 (%)",
        "利润总额累计增长 (%)": "利润总额累计增长 (%)",
        "发电量同比增长 (%)": "发电量同比增长 (%)",
        "工业生产者出厂价格指数 (上年同月=100)": "PPI（上年同月=100）",
        "工业增加值累计增长 (%)": "工业增加值累计增长 (%)",
    }
    wide = wide.rename(columns=rename)
    wide["PPI同比偏离值"] = wide["PPI（上年同月=100）"] - 100
    core = CORE_COLUMNS
    wide["完整核心指标数"] = wide[core].notna().sum(axis=1)
    wide["质量标记"] = wide["完整核心指标数"].eq(len(core)).map({True: "完整", False: "存在官方缺失"})
    wide = wide.reset_index()

    summary = pd.DataFrame({
        "统计项目": ["样本开始月份", "样本结束月份", "月份总数", "完整核心指标月份", "缺失核心指标月份", "生成时间"],
        "结果": [wide["月份"].min(), wide["月份"].max(), len(wide), int((wide["质量标记"] == "完整").sum()), int((wide["质量标记"] != "完整").sum()), datetime.now().strftime("%Y-%m-%d %H:%M")],
    })
    notes = pd.DataFrame({"说明": [
        "数据地区：安徽省；来源：国家统计局国家数据平台分省月度数据。",
        "样本期：2023年08月至2026年07月，共36个月。",
        "PPI同比偏离值 = PPI（上年同月=100） - 100，后续用于 PCA。",
        "工业增加值的1-2月部分月份为空，符合国家统计制度合并发布规则，未进行插值。",
        "累计增长指标保留官方累计口径，未从累计值逆推当月值。",
    ]})
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    core_wide = build_core_wide_table(wide)
    with pd.ExcelWriter(OUTPUT, engine="openpyxl") as writer:
        core_wide.to_excel(writer, sheet_name="核心变量宽表", index=False)
        summary.to_excel(writer, sheet_name="统计摘要", index=False)
        notes.to_excel(writer, sheet_name="数据说明", index=False)
        long.to_excel(writer, sheet_name="原始合并长表", index=False)
    _format_workbook(OUTPUT)
    workbook = load_workbook(OUTPUT)
    workbook["原始合并长表"].sheet_state = "hidden"
    workbook.active = workbook.sheetnames.index("核心变量宽表")
    workbook.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()

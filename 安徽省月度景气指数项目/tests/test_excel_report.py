import numpy as np
import pandas as pd
from openpyxl import load_workbook

from anhui_index.excel_report import write_excel_report
from anhui_index.forecast import forecast_index
from anhui_index.index import fit_prosperity_index


def test_write_excel_report_creates_traceable_result_sheets(tmp_path):
    dates = pd.date_range("2023-01-01", periods=24, freq="MS")
    base = np.arange(24, dtype=float)
    table = pd.DataFrame(
        {
            "date": dates,
            "industrial_yoy": base,
            "investment_cum_yoy": base * 0.8 + 2,
            "profit_cum_yoy": base * 1.2 - 1,
            "power_yoy": base * 0.5 + 4,
            "ppi_yoy": base * 0.3 - 2,
            "quality_flag": "ok",
        }
    )
    index_result = fit_prosperity_index(table)
    forecast_result = forecast_index(index_result.index, horizon=3)
    validation = pd.DataFrame(
        {
            "model": ["sarima"],
            "horizon": [1],
            "mae": [1.0],
            "rmse": [1.2],
            "n_predictions": [5],
        }
    )

    output = write_excel_report(
        table, index_result, forecast_result, validation, tmp_path / "result.xlsx"
    )

    workbook = load_workbook(output, read_only=False, data_only=True)
    assert set(workbook.sheetnames) == {
        "建模月度数据",
        "景气指数",
        "预测结果",
        "滚动验证",
        "主成分载荷",
        "数据说明",
    }
    assert [cell.value for cell in workbook["建模月度数据"][1]] == [
        "时间",
        "变量1：工业增加值同比（%）",
        "变量2：固定资产投资累计同比（%）",
        "变量3：利润总额累计同比（%）",
        "变量4：发电量同比（%）",
        "变量5：PPI同比偏离值（百分点）",
        "质量标记",
    ]
    assert workbook["建模月度数据"]["A2"].number_format == "yyyy-mm-dd"
    assert [cell.value for cell in workbook["预测结果"][1]] == [
        "模型",
        "日期",
        "预测值",
        "下限（95%）",
        "上限（95%）",
        "赤池信息准则（AIC）",
        "贝叶斯信息准则（BIC）",
        "模型参数",
    ]
    assert {row[0] for row in workbook["预测结果"].iter_rows(min_row=2, values_only=True)} == {
        "12个月季节朴素",
        "季节性 ARIMA 模型（SARIMA）",
    }
    assert [cell.value for cell in workbook["滚动验证"][1]] == [
        "模型",
        "预测期限（月）",
        "平均绝对误差（MAE）",
        "均方根误差（RMSE）",
        "预测点数",
    ]
    assert [cell.value for cell in workbook["主成分载荷"][1]] == [
        "指标",
        "主成分分析载荷（PCA）",
    ]
    assert workbook["主成分载荷"][2][0].value == "工业增加值同比（%）"

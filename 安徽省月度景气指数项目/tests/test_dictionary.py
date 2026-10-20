from openpyxl import load_workbook

from anhui_index.dictionary import write_dictionary


def test_write_dictionary_includes_all_core_indicators(tmp_path):
    output = write_dictionary(tmp_path / "data_dictionary.xlsx")

    workbook = load_workbook(output, data_only=True)
    sheet = workbook["指标字典"]
    field_names = {row[0] for row in sheet.iter_rows(min_row=2, values_only=True)}

    assert {
        "industrial_yoy",
        "investment_cum_yoy",
        "profit_cum_yoy",
        "power_yoy",
        "ppi_yoy",
    } <= field_names

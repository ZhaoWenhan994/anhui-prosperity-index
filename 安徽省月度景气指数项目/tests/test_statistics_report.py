from datetime import datetime

from openpyxl import Workbook, load_workbook

from scripts.create_statistics_report import _format_workbook, build_core_wide_table


def test_statistics_report_formats_month_dates_to_fit_column(tmp_path):
    path = tmp_path / "statistics.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["月份"])
    sheet.append([datetime(2023, 8, 1)])
    workbook.save(path)

    _format_workbook(path)

    formatted = load_workbook(path, read_only=False, data_only=False)
    sheet = formatted.active
    assert sheet["A2"].number_format == "yyyy-mm-dd"
    assert sheet.column_dimensions["A"].width >= 12


def test_core_wide_table_has_time_rows_and_five_numbered_variables():
    source = __import__("pandas").DataFrame(
        {
            "月份": __import__("pandas").date_range("2023-01-01", periods=2, freq="MS"),
            "工业增加值同比增长 (%)": [1.0, 2.0],
            "固定资产投资累计增长 (%)": [3.0, 4.0],
            "利润总额累计增长 (%)": [5.0, 6.0],
            "发电量同比增长 (%)": [7.0, 8.0],
            "PPI同比偏离值": [9.0, 10.0],
            "完整核心指标数": [5, 5],
            "质量标记": ["完整", "完整"],
        }
    )

    wide = build_core_wide_table(source)

    assert list(wide.columns) == [
        "时间",
        "变量1：工业增加值同比（%）",
        "变量2：固定资产投资累计同比（%）",
        "变量3：利润总额累计同比（%）",
        "变量4：发电量同比（%）",
        "变量5：PPI同比偏离值（百分点）",
        "完整核心指标数",
        "质量标记",
    ]
    assert len(wide) == 2

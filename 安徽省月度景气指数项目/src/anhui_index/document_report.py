from __future__ import annotations

from datetime import date, datetime
from html import escape
import os
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from anhui_index.excel_report import (
    INDICATOR_LABELS,
    localize_forecast_table,
    localize_validation_table,
)
from anhui_index.forecast import ForecastResult
from anhui_index.index import IndexResult


REPORT_TITLE = "安徽省月度生产与投资景气指数分析报告"


def _display_value(value: object) -> str:
    if value is None or (not isinstance(value, (list, tuple, dict)) and pd.isna(value)):
        return ""
    if isinstance(value, (pd.Timestamp, datetime, date)):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float):
        return f"{value:.4f}"
    return str(value)


def _localized_tables(
    index_result: IndexResult,
    forecast_result: ForecastResult,
    validation: pd.DataFrame,
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    loading = (
        index_result.loadings.rename("主成分分析载荷（PCA）")
        .rename(index=INDICATOR_LABELS)
        .rename_axis("指标")
        .reset_index()
    )
    forecast = localize_forecast_table(forecast_result.forecasts)
    forecast = forecast[["模型", "日期", "预测值", "下限（95%）", "上限（95%）"]]
    validation_display = localize_validation_table(validation)
    return loading, forecast, validation_display


def _set_docx_font(document: Document) -> None:
    for style_name in ("Normal", "Title", "Heading 1", "Heading 2"):
        style = document.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style.font.size = Pt(10.5 if style_name == "Normal" else 14)
        style.font.bold = style_name in ("Title", "Heading 1", "Heading 2")
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def _set_run_font(run, bold: bool) -> None:
    run.font.name = "Microsoft YaHei"
    run.font.bold = bold
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    for font_attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn(f"w:{font_attribute}"), "Microsoft YaHei")


def _apply_direct_run_format(document: Document) -> None:
    for paragraph in document.paragraphs:
        bold = paragraph.style.style_id in {"Title", "Heading1", "Heading2"}
        for run in paragraph.runs:
            _set_run_font(run, bold=bold)
    for table in document.tables:
        for row_index, row in enumerate(table.rows):
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        _set_run_font(run, bold=row_index == 0)


def _add_docx_table(document: Document, frame: pd.DataFrame) -> None:
    table = document.add_table(rows=1, cols=len(frame.columns))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, label in zip(table.rows[0].cells, frame.columns):
        cell.text = str(label)
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for row in frame.itertuples(index=False, name=None):
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = _display_value(value)
    document.add_paragraph()


def _add_docx_image(document: Document, path: Path) -> None:
    if path.exists():
        document.add_picture(str(path), width=Inches(6.2))
        document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER


def generate_word_report(
    table: pd.DataFrame,
    index_result: IndexResult,
    forecast_result: ForecastResult,
    validation: pd.DataFrame,
    output_dir: Path,
) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    figures = output_dir / "figures"
    loading, forecast, validation_display = _localized_tables(
        index_result, forecast_result, validation
    )
    document = Document()
    _set_docx_font(document)
    title = document.add_heading(REPORT_TITLE, level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph(
        "本报告使用国家统计局分省月度数据，采用主成分分析（PCA）构建安徽省月度景气指数，"
        "并使用季节性 ARIMA 模型（SARIMA）和12个月季节朴素方法进行短期预测。"
    )

    document.add_heading("一、样本与方法", level=1)
    for text in (
        f"样本期：{pd.to_datetime(table['date']).min():%Y-%m-%d} 至 {pd.to_datetime(table['date']).max():%Y-%m-%d}，共 {len(table)} 个月。",
        f"景气指数观测月：{index_result.index.index.min():%Y-%m-%d} 至 {index_result.index.index.max():%Y-%m-%d}。",
        "核心指标：工业增加值同比、固定资产投资累计同比、利润总额累计同比、发电量同比、工业生产者出厂价格指数（PPI）同比偏离值。",
        f"指数基准：{index_result.baseline_year}年平均值缩放为100，使用{index_result.baseline_observations}个月观测。",
        "固定资产投资和利润使用官方累计增速，不能直接解释为当月增速；官方缺失月份未插值。",
    ):
        document.add_paragraph(text, style="List Bullet")

    document.add_heading("二、核心指标走势", level=1)
    _add_docx_image(document, figures / "indicator_trends.png")
    document.add_heading("三、主成分分析载荷（PCA）", level=1)
    _add_docx_image(document, figures / "pca_loadings.png")
    _add_docx_table(document, loading)
    document.add_heading("四、景气指数与预测", level=1)
    _add_docx_image(document, figures / "index_forecast.png")
    _add_docx_table(document, forecast)
    document.add_heading("五、滚动验证", level=1)
    document.add_paragraph("平均绝对误差（MAE）和均方根误差（RMSE）越小，表示预测误差越小。")
    _add_docx_table(document, validation_display)
    document.add_heading("六、解释限制", level=1)
    document.add_paragraph(
        "当前样本长度较短，预测结果适合作为应用统计复试项目中的方法展示和短期参考，"
        "不应视为对未来经济走势的确定性判断。"
    )
    _apply_direct_run_format(document)

    path = output_dir / "安徽省月度景气指数分析报告.docx"
    document.save(path)
    return path


def _register_chinese_pdf_font() -> str:
    font_name = "AnhuiChinese"
    if font_name in pdfmetrics.getRegisteredFontNames():
        return font_name
    candidates = (
        Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts" / "simhei.ttf",
        Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts" / "NotoSansSC-VF.ttf",
    )
    for path in candidates:
        if path.exists():
            pdfmetrics.registerFont(TTFont(font_name, str(path)))
            return font_name
    return "Helvetica"


def _pdf_table(frame: pd.DataFrame, font_name: str) -> Table:
    header_style = ParagraphStyle(
        "TableHeader", fontName=font_name, fontSize=8, leading=10, alignment=TA_CENTER
    )
    cell_style = ParagraphStyle("TableCell", fontName=font_name, fontSize=8, leading=10)
    data = [[Paragraph(escape(str(column)), header_style) for column in frame.columns]]
    for row in frame.itertuples(index=False, name=None):
        data.append([Paragraph(escape(_display_value(value)), cell_style) for value in row])
    table = Table(data, repeatRows=1, hAlign="LEFT")
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#176B87")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#B8CDD4")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    return table


def _pdf_image(path: Path) -> Image | None:
    if not path.exists():
        return None
    image = Image(str(path))
    max_width, max_height = 175 * mm, 100 * mm
    scale = min(max_width / image.drawWidth, max_height / image.drawHeight, 1)
    image.drawWidth *= scale
    image.drawHeight *= scale
    return image


def generate_pdf_report(
    table: pd.DataFrame,
    index_result: IndexResult,
    forecast_result: ForecastResult,
    validation: pd.DataFrame,
    output_dir: Path,
) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    figures = output_dir / "figures"
    loading, forecast, validation_display = _localized_tables(
        index_result, forecast_result, validation
    )
    font_name = _register_chinese_pdf_font()
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "ReportTitle", parent=styles["Title"], fontName=font_name, fontSize=18,
        leading=23, alignment=TA_CENTER, textColor=colors.HexColor("#176B87")
    )
    body_style = ParagraphStyle(
        "ReportBody", parent=styles["BodyText"], fontName=font_name, fontSize=10,
        leading=16
    )
    heading_style = ParagraphStyle(
        "ReportHeading", parent=styles["Heading2"], fontName=font_name, fontSize=13,
        leading=18, textColor=colors.HexColor("#176B87")
    )
    document = SimpleDocTemplate(
        str(output_dir / "安徽省月度景气指数分析报告.pdf"),
        pagesize=A4,
        rightMargin=18 * mm,
        leftMargin=18 * mm,
        topMargin=16 * mm,
        bottomMargin=16 * mm,
    )
    story = [
        Paragraph(REPORT_TITLE, title_style),
        Spacer(1, 8),
        Paragraph(
            "本报告使用国家统计局分省月度数据，采用主成分分析（PCA）构建安徽省月度景气指数，"
            "并使用季节性 ARIMA 模型（SARIMA）和12个月季节朴素方法进行短期预测。",
            body_style,
        ),
        Spacer(1, 8),
        Paragraph("一、样本与方法", heading_style),
        Paragraph(
            f"样本期：{pd.to_datetime(table['date']).min():%Y-%m-%d} 至 {pd.to_datetime(table['date']).max():%Y-%m-%d}，共 {len(table)} 个月。<br/>"
            f"指数基准：{index_result.baseline_year}年平均值缩放为100，使用{index_result.baseline_observations}个月观测。<br/>"
            "核心指标包括工业增加值同比、固定资产投资累计同比、利润总额累计同比、发电量同比和工业生产者出厂价格指数（PPI）同比偏离值。<br/>"
            "固定资产投资和利润使用官方累计增速，官方缺失月份未插值。",
            body_style,
        ),
        Spacer(1, 8),
        Paragraph("二、核心指标走势", heading_style),
    ]
    for figure_name in ("indicator_trends.png", "pca_loadings.png", "index_forecast.png"):
        image = _pdf_image(figures / figure_name)
        if image is not None:
            story.extend([image, Spacer(1, 8)])
    story.extend(
        [
            Paragraph("三、主成分分析载荷（PCA）", heading_style),
            _pdf_table(loading, font_name),
            Spacer(1, 10),
            Paragraph("四、预测结果", heading_style),
            _pdf_table(forecast, font_name),
            Spacer(1, 10),
            Paragraph("五、滚动验证", heading_style),
            Paragraph("平均绝对误差（MAE）和均方根误差（RMSE）越小，表示预测误差越小。", body_style),
            _pdf_table(validation_display, font_name),
            Spacer(1, 10),
            Paragraph("六、解释限制", heading_style),
            Paragraph(
                "当前样本长度较短，预测结果适合作为应用统计复试项目中的方法展示和短期参考，"
                "不应视为对未来经济走势的确定性判断。",
                body_style,
            ),
        ]
    )
    path = output_dir / "安徽省月度景气指数分析报告.pdf"
    document.build(story)
    return path


def generate_document_reports(
    table: pd.DataFrame,
    index_result: IndexResult,
    forecast_result: ForecastResult,
    validation: pd.DataFrame,
    output_dir: Path,
) -> tuple[Path, Path]:
    """Generate editable Word and printable PDF versions of the report."""
    word_path = generate_word_report(
        table, index_result, forecast_result, validation, output_dir
    )
    pdf_path = generate_pdf_report(
        table, index_result, forecast_result, validation, output_dir
    )
    return word_path, pdf_path

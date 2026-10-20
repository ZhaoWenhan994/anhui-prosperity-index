from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from anhui_index.forecast import ForecastResult
from anhui_index.index import IndexResult


INDICATOR_LABELS = {
    "industrial_yoy": "工业增加值同比（%）",
    "investment_cum_yoy": "固定资产投资累计同比（%）",
    "profit_cum_yoy": "利润总额累计同比（%）",
    "power_yoy": "发电量同比（%）",
    "ppi_yoy": "工业生产者出厂价格指数（PPI）同比偏离值（百分点）",
}

VARIABLE_LABELS = {
    "industrial_yoy": "变量1：工业增加值同比（%）",
    "investment_cum_yoy": "变量2：固定资产投资累计同比（%）",
    "profit_cum_yoy": "变量3：利润总额累计同比（%）",
    "power_yoy": "变量4：发电量同比（%）",
    "ppi_yoy": "变量5：PPI同比偏离值（百分点）",
}

TABLE_COLUMN_LABELS = {
    "date": "时间",
    "industrial_yoy": VARIABLE_LABELS["industrial_yoy"],
    "investment_cum_yoy": VARIABLE_LABELS["investment_cum_yoy"],
    "profit_cum_yoy": VARIABLE_LABELS["profit_cum_yoy"],
    "power_yoy": VARIABLE_LABELS["power_yoy"],
    "ppi_yoy": VARIABLE_LABELS["ppi_yoy"],
    "quality_flag": "质量标记",
}

MODEL_LABELS = {
    "sarima": "季节性 ARIMA 模型（SARIMA）",
    "seasonal_naive": "12个月季节朴素",
}

FORECAST_COLUMN_LABELS = {
    "model": "模型",
    "date": "日期",
    "forecast": "预测值",
    "lower_95": "下限（95%）",
    "upper_95": "上限（95%）",
    "aic": "赤池信息准则（AIC）",
    "bic": "贝叶斯信息准则（BIC）",
    "model_detail": "模型参数",
}

VALIDATION_COLUMN_LABELS = {
    "model": "模型",
    "horizon": "预测期限（月）",
    "mae": "平均绝对误差（MAE）",
    "rmse": "均方根误差（RMSE）",
    "n_predictions": "预测点数",
}


def _localize_model_detail(value: object) -> object:
    if value == "t-12":
        return "滞后12个月"
    if isinstance(value, str):
        return value.replace("seasonal_order=", "季节阶数=").replace(
            "order=", "非季节阶数="
        )
    return value


def _localize_table(table: pd.DataFrame) -> pd.DataFrame:
    localized = table.rename(columns=TABLE_COLUMN_LABELS).copy()
    if "质量标记" in localized:
        localized["质量标记"] = localized["质量标记"].map(
            {"ok": "完整", "missing_core_indicator": "官方缺失"}
        ).fillna(localized["质量标记"])
    return localized


def localize_index_series(index: pd.Series) -> pd.DataFrame:
    """Convert the internal index series to a Chinese result table."""
    return index.rename("景气指数").rename_axis("日期").reset_index()


def localize_forecast_table(table: pd.DataFrame) -> pd.DataFrame:
    """Convert internal forecast fields to Chinese display labels."""
    localized = table.rename(columns=FORECAST_COLUMN_LABELS).copy()
    localized["模型"] = localized["模型"].map(MODEL_LABELS).fillna(localized["模型"])
    localized["模型参数"] = localized["模型参数"].map(_localize_model_detail)
    return localized


def localize_validation_table(table: pd.DataFrame) -> pd.DataFrame:
    """Convert internal validation fields to Chinese display labels."""
    localized = table.rename(columns=VALIDATION_COLUMN_LABELS).copy()
    localized["模型"] = localized["模型"].map(MODEL_LABELS).fillna(localized["模型"])
    return localized


def _format_workbook(path: Path) -> None:
    workbook = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="176B87")
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, (date, datetime)):
                    cell.number_format = "yyyy-mm-dd"
        for cell in sheet[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = header_fill
        for column_cells in sheet.columns:
            values = [str(cell.value) if cell.value is not None else "" for cell in column_cells]
            width = min(max(len(value) for value in values) + 2, 42)
            sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = width
    workbook.save(path)


def write_excel_report(
    table: pd.DataFrame,
    index_result: IndexResult,
    forecast_result: ForecastResult,
    validation: pd.DataFrame,
    path: Path,
) -> Path:
    """Write modelling inputs and results to a traceable Excel workbook."""
    path.parent.mkdir(parents=True, exist_ok=True)
    localized_table = _localize_table(table)
    index_table = localize_index_series(index_result.index)
    loading_table = (
        index_result.loadings.rename("主成分分析载荷（PCA）")
        .rename_axis("指标")
        .reset_index()
    )
    loading_table["指标"] = loading_table["指标"].map(INDICATOR_LABELS)
    forecast_table = localize_forecast_table(forecast_result.forecasts)
    validation_table = localize_validation_table(validation)
    notes = pd.DataFrame(
        {
            "说明": [
                "数据地区：安徽省；来源：国家统计局国家数据平台分省月度数据。",
                "主成分分析（PCA）输入为五项核心指标，只有核心指标齐全的月份进入指数。",
                "工业生产者出厂价格指数（PPI）已由上年同月=100转换为数值减100的同比偏离值。",
                "固定资产投资和利润使用官方累计增速，不能直接解释为当月增速。",
                "官方缺失月份未插值；当前原始表未提供 CPI 和季度社会零售数据。",
                f"指数基准年：{index_result.baseline_year}年，使用 {index_result.baseline_observations} 个月观测。",
            ]
        }
    )
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        localized_table.to_excel(writer, sheet_name="建模月度数据", index=False)
        index_table.to_excel(writer, sheet_name="景气指数", index=False)
        forecast_table.to_excel(writer, sheet_name="预测结果", index=False)
        validation_table.to_excel(writer, sheet_name="滚动验证", index=False)
        loading_table.to_excel(writer, sheet_name="主成分载荷", index=False)
        notes.to_excel(writer, sheet_name="数据说明", index=False)
    _format_workbook(path)
    return path

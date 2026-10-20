from __future__ import annotations

import re
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from anhui_index.metadata import INDICATORS


CANONICAL_COLUMNS = ["date", "indicator", "value", "source_file", "official_name"]
_MONTH_PATTERN = re.compile(r"^(\d{4})年(\d{1,2})月$")
_OFFICIAL_NAMES = {spec.key: spec.official_name for spec in INDICATORS}


def _normalize(value: object) -> str:
    return str(value).replace("\u3000", " ").strip()


def _parse_month(value: object, path: Path) -> str:
    match = _MONTH_PATTERN.match(_normalize(value))
    if not match:
        raise ValueError(f"{path.name} 中存在无法识别的月份表头：{value!r}")
    return f"{match.group(1)}-{int(match.group(2)):02d}-01"


def import_workbook(path: Path, indicator_key: str) -> pd.DataFrame:
    """Parse one Statistical Bureau export into canonical observed-value records."""
    if indicator_key not in _OFFICIAL_NAMES:
        raise ValueError(f"未知指标字段：{indicator_key}")

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if len(rows) < 5:
        raise ValueError(f"{path.name} 不符合国家统计局月度导出表的最小行数要求")

    headers = list(rows[3])
    dates = [_parse_month(value, path) for value in headers[1:] if value is not None]
    expected_name = _OFFICIAL_NAMES[indicator_key]
    indicator_row = next(
        (row for row in rows[4:] if row and _normalize(row[0]) == expected_name), None
    )
    if indicator_row is None:
        raise ValueError(f"{path.name} 中未找到预期官方指标名称：{expected_name}")

    records: list[dict[str, object]] = []
    for date, value in zip(dates, indicator_row[1:]):
        if value in (None, ""):
            continue
        records.append(
            {
                "date": date,
                "indicator": indicator_key,
                "value": float(value),
                "source_file": path.name,
                "official_name": expected_name,
            }
        )

    return pd.DataFrame(records, columns=CANONICAL_COLUMNS).sort_values(
        "date", ignore_index=True
    )


def import_manifest(manifest_path: Path, raw_directory: Path) -> pd.DataFrame:
    """Import every indicator specified by a CSV manifest."""
    manifest = pd.read_csv(manifest_path)
    required = {"indicator", "filename"}
    if not required <= set(manifest.columns):
        raise ValueError("原始文件清单必须包含 indicator 和 filename 两列")

    tables = []
    for entry in manifest.itertuples(index=False):
        path = raw_directory / entry.filename
        if not path.exists():
            raise FileNotFoundError(f"缺少原始数据文件：{path.name}")
        tables.append(import_workbook(path, entry.indicator))
    return pd.concat(tables, ignore_index=True) if tables else pd.DataFrame(columns=CANONICAL_COLUMNS)
